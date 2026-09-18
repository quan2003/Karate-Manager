from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "Tai_lieu"
OUT_FILE = OUT_DIR / "Huong_dan_quy_trinh_check_the_va_can_VDV_2026.docx"
LOGO = ROOT / "src" / "assets" / "icon.png"

NAVY = "173B67"
BLUE = "2563EB"
LIGHT_BLUE = "EAF2FF"
LIGHT_GRAY = "F4F7FA"
MID_GRAY = "D7E1EC"
TEXT = "172033"
MUTED = "52677F"
YELLOW = "FFF7D6"
YELLOW_BORDER = "E5A900"
RED = "B42318"
LIGHT_RED = "FFF1F0"
GREEN = "16803C"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name, edge_data in edges.items():
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        for key, value in edge_data.items():
            edge.set(qn(f"w:{key}"), str(value))


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Times New Roman", size=12, bold=False, color=TEXT, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, align=None, before=0, after=4, line=1.12, keep_with_next=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep_with_next
    if align is not None:
        paragraph.alignment = align


def add_rich_paragraph(doc, parts, *, align=None, before=0, after=4, line=1.12, indent=0, keep=False):
    p = doc.add_paragraph()
    style_paragraph(p, align, before, after, line, keep)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
        p.paragraph_format.first_line_indent = Cm(-0.45)
    for part in parts:
        if isinstance(part, str):
            text, bold, color, italic = part, False, TEXT, False
        else:
            text = part.get("text", "")
            bold = part.get("bold", False)
            color = part.get("color", TEXT)
            italic = part.get("italic", False)
        run = p.add_run(text)
        set_run_font(run, bold=bold, color=color, italic=italic)
    return p


def add_bullet(doc, text, *, bold_prefix=None, color=TEXT):
    p = doc.add_paragraph()
    style_paragraph(p, after=2, line=1.08)
    p.paragraph_format.left_indent = Cm(0.62)
    p.paragraph_format.first_line_indent = Cm(-0.38)
    bullet = p.add_run("•  ")
    set_run_font(bullet, bold=True, color=BLUE)
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, bold=True, color=color)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest, color=color)
    else:
        run = p.add_run(text)
        set_run_font(run, color=color)
    return p


def add_section_heading(doc, number, title):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(0.9)
    table.columns[1].width = Cm(15.8)
    table.rows[0].cells[0].width = Cm(0.9)
    table.rows[0].cells[1].width = Cm(15.8)
    for cell in table.rows[0].cells:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell, top=70, bottom=70, start=80, end=80)
        set_cell_border(cell, bottom={"val": "single", "sz": "10", "color": BLUE})
    set_cell_shading(table.cell(0, 0), BLUE)
    p_num = table.cell(0, 0).paragraphs[0]
    style_paragraph(p_num, WD_ALIGN_PARAGRAPH.CENTER, after=0)
    r_num = p_num.add_run(str(number))
    set_run_font(r_num, size=11.5, bold=True, color="FFFFFF")
    p_title = table.cell(0, 1).paragraphs[0]
    style_paragraph(p_title, after=0, keep_with_next=True)
    r_title = p_title.add_run(title.upper())
    set_run_font(r_title, size=12.5, bold=True, color=NAVY)
    spacer = doc.add_paragraph()
    style_paragraph(spacer, after=1)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_cell_text(cell, text, *, bold=False, color=TEXT, align=WD_ALIGN_PARAGRAPH.LEFT, size=11):
    p = cell.paragraphs[0]
    style_paragraph(p, align, after=0, line=1.08)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell, top=100, bottom=100, start=110, end=110)
    return p


def add_note_box(doc, title, bullets, *, fill=YELLOW, border=YELLOW_BORDER, title_color=RED):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=150, bottom=140, start=170, end=170)
    edge = {"val": "single", "sz": "12", "color": border}
    set_cell_border(cell, top=edge, left=edge, bottom=edge, right=edge)
    p = cell.paragraphs[0]
    style_paragraph(p, after=5, keep_with_next=True)
    r = p.add_run(title)
    set_run_font(r, size=12, bold=True, color=title_color)
    for item in bullets:
        p = cell.add_paragraph()
        style_paragraph(p, after=2, line=1.08)
        p.paragraph_format.left_indent = Cm(0.42)
        p.paragraph_format.first_line_indent = Cm(-0.32)
        marker = p.add_run("•  ")
        set_run_font(marker, size=11, bold=True, color=title_color)
        body = p.add_run(item)
        set_run_font(body, size=11, color=TEXT)
    after = doc.add_paragraph()
    style_paragraph(after, after=1)


def build_document():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    section.header_distance = Cm(0.55)
    section.footer_distance = Cm(0.65)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor.from_string(TEXT)

    # Header identity block
    header_table = doc.add_table(rows=1, cols=3)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    widths = [Cm(2.1), Cm(10.8), Cm(3.8)]
    for idx, width in enumerate(widths):
        header_table.columns[idx].width = width
        header_table.cell(0, idx).width = width
        set_cell_margins(header_table.cell(0, idx), top=30, bottom=80, start=70, end=70)
        set_cell_border(header_table.cell(0, idx), bottom={"val": "single", "sz": "12", "color": NAVY})
        header_table.cell(0, idx).vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if LOGO.exists():
        p_logo = header_table.cell(0, 0).paragraphs[0]
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.add_run().add_picture(str(LOGO), width=Cm(1.45))
    p_center = header_table.cell(0, 1).paragraphs[0]
    style_paragraph(p_center, WD_ALIGN_PARAGRAPH.CENTER, after=1)
    r = p_center.add_run("GIẢI CÁC CÂU LẠC BỘ NGHĨA DŨNG KARATE-DO\nPHÂN ĐƯỜNG ĐÀ NẴNG 3 LẦN THỨ I - NĂM 2026")
    set_run_font(r, size=10.5, bold=True, color=NAVY)
    p_right = header_table.cell(0, 2).paragraphs[0]
    style_paragraph(p_right, WD_ALIGN_PARAGRAPH.CENTER, after=1)
    r = p_right.add_run("BAN TỔ CHỨC\nHƯỚNG DẪN NỘI BỘ")
    set_run_font(r, size=9.5, bold=True, color=NAVY)

    p_title = doc.add_paragraph()
    style_paragraph(p_title, WD_ALIGN_PARAGRAPH.CENTER, before=12, after=4, line=1.05, keep_with_next=True)
    r = p_title.add_run("HƯỚNG DẪN QUY TRÌNH")
    set_run_font(r, size=16, bold=True, color=NAVY)
    p_subtitle = doc.add_paragraph()
    style_paragraph(p_subtitle, WD_ALIGN_PARAGRAPH.CENTER, after=10, line=1.05, keep_with_next=True)
    r = p_subtitle.add_run("KIỂM TRA THẺ VÀ CÂN VẬN ĐỘNG VIÊN")
    set_run_font(r, size=15, bold=True, color=BLUE)

    info = doc.add_table(rows=2, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.autofit = False
    info.columns[0].width = Cm(8.35)
    info.columns[1].width = Cm(8.35)
    labels = [
        ("Thời gian", "........................................................"),
        ("Địa điểm", "........................................................"),
        ("Thời hạn hoàn tất", "................................................"),
        ("Dung sai áp dụng", "±2 kg"),
    ]
    for index, (label, value) in enumerate(labels):
        cell = info.cell(index // 2, index % 2)
        set_cell_shading(cell, LIGHT_BLUE if index == 3 else LIGHT_GRAY)
        edge = {"val": "single", "sz": "6", "color": MID_GRAY}
        set_cell_border(cell, top=edge, left=edge, bottom=edge, right=edge)
        p = cell.paragraphs[0]
        style_paragraph(p, after=0)
        r1 = p.add_run(f"{label}: ")
        set_run_font(r1, size=10.5, bold=True, color=NAVY)
        r2 = p.add_run(value)
        set_run_font(r2, size=10.5, bold=index == 3, color=BLUE if index == 3 else TEXT)
        set_cell_margins(cell, top=95, bottom=95, start=120, end=120)

    add_rich_paragraph(
        doc,
        [
            {"text": "Kính gửi: ", "bold": True, "color": NAVY},
            "Huấn luyện viên, Trưởng đoàn các Câu lạc bộ/Đơn vị tham dự giải.",
        ],
        before=9,
        after=4,
    )
    add_rich_paragraph(
        doc,
        [
            "Nhằm bảo đảm công tác kiểm tra VĐV được thực hiện ",
            {"text": "nhanh chóng, chính xác và minh bạch", "bold": True},
            ", Ban Tổ chức hướng dẫn quy trình kiểm tra thẻ và cân VĐV như sau:",
        ],
        after=8,
    )

    add_section_heading(doc, 1, "Đối tượng và nguyên tắc áp dụng")
    add_bullet(doc, "Tất cả VĐV tham dự giải phải thực hiện kiểm tra thẻ VĐV.")
    add_bullet(doc, "VĐV thi đấu Kata chỉ kiểm tra thẻ, không thực hiện cân.")
    add_bullet(doc, "VĐV thi đấu Kumite phải kiểm tra thẻ và cân trọng lượng thực tế.")
    add_bullet(doc, "Check thẻ được dùng chung cho một VĐV; số cân được ghi nhận riêng theo từng nội dung Kumite đã đăng ký.")
    add_bullet(doc, "Kết quả cân được đối chiếu với hạng cân của nội dung thi đấu, không dựa vào số cân VĐV tự khai khi đăng ký.")

    add_section_heading(doc, 2, "Hồ sơ và công tác chuẩn bị")
    add_bullet(doc, "Thẻ VĐV do Ban Tổ chức phát hành.")
    add_bullet(doc, "Giấy tờ tùy thân hoặc giấy tờ xác minh độ tuổi khi Ban Tổ chức yêu cầu.")
    add_bullet(doc, "Danh sách VĐV của đoàn theo từng nội dung thi đấu.")
    add_bullet(doc, "HLV/Trưởng đoàn hoặc người được ủy quyền đi cùng để phối hợp xác nhận khi có sai lệch.")

    add_section_heading(doc, 3, "Quy trình thực hiện tại bàn Check-in")
    process = doc.add_table(rows=1, cols=3)
    process.alignment = WD_TABLE_ALIGNMENT.CENTER
    process.autofit = False
    process_widths = [Cm(1.2), Cm(4.0), Cm(11.5)]
    headers = ["BƯỚC", "NỘI DUNG", "CÁCH THỰC HIỆN"]
    for idx, text in enumerate(headers):
        process.columns[idx].width = process_widths[idx]
        cell = process.cell(0, idx)
        cell.width = process_widths[idx]
        set_cell_shading(cell, NAVY)
        add_cell_text(cell, text, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5)
    set_repeat_table_header(process.rows[0])
    steps = [
        ("1", "Tiếp nhận VĐV", "VĐV đến đúng khung giờ, xếp hàng theo hướng dẫn. Đại diện đoàn chuẩn bị sẵn thẻ và danh sách để đối chiếu."),
        ("2", "Kiểm tra thẻ", "Thư ký đối chiếu họ tên, năm sinh, đơn vị, nội dung đăng ký và giấy tờ liên quan. Hợp lệ thì xác nhận “Đã check thẻ” trên hệ thống."),
        ("3", "Cân VĐV Kumite", "VĐV lên cân theo hướng dẫn. Thư ký nhập số cân thực tế; hệ thống tự đối chiếu với hạng cân và mức dung sai. Vượt giới hạn sẽ hiển thị cảnh báo màu đỏ."),
        ("4", "Rà soát kết quả", "HLV/Trưởng đoàn kiểm tra lại trạng thái của đoàn. Các trường hợp chưa hoàn tất hoặc có cảnh báo phải được xử lý trước thời hạn quy định."),
    ]
    for row_index, values in enumerate(steps, start=1):
        cells = process.add_row().cells
        for col_index, text in enumerate(values):
            cells[col_index].width = process_widths[col_index]
            set_cell_shading(cells[col_index], "FFFFFF" if row_index % 2 else LIGHT_GRAY)
            edge = {"val": "single", "sz": "5", "color": MID_GRAY}
            set_cell_border(cells[col_index], top=edge, left=edge, bottom=edge, right=edge)
            add_cell_text(
                cells[col_index],
                text,
                bold=col_index in (0, 1),
                color=BLUE if col_index == 0 else NAVY if col_index == 1 else TEXT,
                align=WD_ALIGN_PARAGRAPH.CENTER if col_index == 0 else WD_ALIGN_PARAGRAPH.LEFT,
                size=10.5,
            )

    add_rich_paragraph(doc, [{"text": "Ví dụ đối chiếu hạng cân", "bold": True, "color": NAVY}], before=8, after=3, keep=True)
    example = doc.add_table(rows=1, cols=3)
    example.alignment = WD_TABLE_ALIGNMENT.CENTER
    example.autofit = False
    example_values = [
        ("Hạng cân nội dung", "35 kg"),
        ("Dung sai", "±2 kg"),
        ("Khoảng hợp lệ", "33 - 37 kg"),
    ]
    for idx, (label, value) in enumerate(example_values):
        cell = example.cell(0, idx)
        cell.width = Cm(5.55)
        set_cell_shading(cell, LIGHT_BLUE)
        edge = {"val": "single", "sz": "8", "color": "9FC0ED"}
        set_cell_border(cell, top=edge, left=edge, bottom=edge, right=edge)
        p = cell.paragraphs[0]
        style_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER, after=2)
        r1 = p.add_run(label + "\n")
        set_run_font(r1, size=9.5, bold=True, color=MUTED)
        r2 = p.add_run(value)
        set_run_font(r2, size=13, bold=True, color=BLUE)
        set_cell_margins(cell, top=110, bottom=110, start=90, end=90)

    add_note_box(
        doc,
        "LƯU Ý QUAN TRỌNG",
        [
            "Kata chỉ check thẻ; không nhập cân cho VĐV chỉ tham gia Kata.",
            "Kumite phải hoàn tất cả check thẻ và cân thực tế mới được xem là hoàn tất Check-in.",
            "Mốc cân được lấy từ nội dung thi đấu. Dung sai ±2 kg được áp dụng ở cả hai phía của mốc hoặc khoảng cân.",
            "HLV/Trưởng đoàn phải kiểm tra kết quả trước khi rời bàn. Khiếu nại về số cân cần thực hiện ngay tại khu vực cân.",
            "Không tự ý yêu cầu thư ký sửa hạng mục, chuyển hạng cân hoặc thay đổi thông tin đăng ký tại bàn Check-in.",
            "Khi hệ thống hoặc mạng LAN gặp sự cố, sử dụng danh sách giấy dự phòng và có chữ ký xác nhận của đại diện đoàn.",
        ],
    )

    add_section_heading(doc, 4, "Xử lý các trường hợp phát sinh")
    issues = doc.add_table(rows=1, cols=2)
    issues.alignment = WD_TABLE_ALIGNMENT.CENTER
    issues.autofit = False
    issue_widths = [Cm(5.0), Cm(11.7)]
    for idx, text in enumerate(("TRƯỜNG HỢP", "CÁCH XỬ LÝ")):
        issues.columns[idx].width = issue_widths[idx]
        cell = issues.cell(0, idx)
        cell.width = issue_widths[idx]
        set_cell_shading(cell, NAVY)
        add_cell_text(cell, text, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5)
    set_repeat_table_header(issues.rows[0])
    issue_rows = [
        ("Sai họ tên, năm sinh, đơn vị", "Tạm dừng xác nhận; chuyển bộ phận chuyên môn/Ban Tổ chức kiểm tra và xử lý."),
        ("Quên hoặc mất thẻ VĐV", "Báo Ban Tổ chức để xác minh. Thư ký không tự ý xác nhận khi chưa đủ căn cứ."),
        ("Không đạt hạng cân", "Thông báo ngay cho VĐV và đại diện đoàn. Việc cân lại hoặc xử lý tiếp theo thực hiện theo quyết định của Ban Tổ chức."),
        ("Đăng ký sai nội dung", "Không sửa trực tiếp tại bàn cân. HLV/Trưởng đoàn liên hệ Ban Tổ chức để được xem xét."),
        ("Mất mạng/lỗi thiết bị", "Ghi nhận trên danh sách giấy dự phòng, có chữ ký xác nhận; cập nhật lại hệ thống khi hoạt động ổn định."),
    ]
    for row_index, values in enumerate(issue_rows, start=1):
        cells = issues.add_row().cells
        for col_index, text in enumerate(values):
            cells[col_index].width = issue_widths[col_index]
            set_cell_shading(cells[col_index], "FFFFFF" if row_index % 2 else LIGHT_GRAY)
            edge = {"val": "single", "sz": "5", "color": MID_GRAY}
            set_cell_border(cells[col_index], top=edge, left=edge, bottom=edge, right=edge)
            add_cell_text(cells[col_index], text, bold=col_index == 0, color=NAVY if col_index == 0 else TEXT, size=10.5)

    add_section_heading(doc, 5, "Trách nhiệm của HLV/Trưởng đoàn")
    responsibilities = [
        "Kiểm tra trước thông tin đăng ký của toàn bộ VĐV thuộc đơn vị.",
        "Phổ biến lịch, địa điểm và quy trình Check-in cho VĐV.",
        "Bố trí VĐV đến đúng giờ và cử người đại diện phối hợp xử lý sai lệch.",
        "Rà soát kết quả Check-in của đoàn trước thời hạn Ban Tổ chức quy định.",
        "Ký xác nhận đối với trường hợp sai thông tin, không đạt hạng cân, sử dụng phiếu giấy dự phòng hoặc phát sinh khiếu nại.",
    ]
    for item in responsibilities:
        add_bullet(doc, item)

    add_note_box(
        doc,
        "HOÀN TẤT CHECK-IN KHI",
        [
            "VĐV Kata: trạng thái thẻ đã được xác nhận.",
            "VĐV Kumite: thẻ đã được xác nhận, đã nhập cân thực tế và không có cảnh báo chưa xử lý.",
            "HLV/Trưởng đoàn đã rà soát toàn bộ danh sách của đơn vị.",
        ],
        fill=LIGHT_RED,
        border="E7A19A",
        title_color=RED,
    )

    closing = add_rich_paragraph(
        doc,
        [
            "Ban Tổ chức đề nghị HLV/Trưởng đoàn các đơn vị phối hợp thực hiện nghiêm túc để công tác kiểm tra VĐV diễn ra đúng tiến độ, chính xác và minh bạch.",
        ],
        before=4,
        after=12,
        line=1.12,
    )
    closing.paragraph_format.keep_together = True

    signatures = doc.add_table(rows=1, cols=2)
    signatures.alignment = WD_TABLE_ALIGNMENT.CENTER
    signatures.autofit = False
    for idx, title in enumerate(("ĐẠI DIỆN BỘ PHẬN CHECK-IN", "BAN TỔ CHỨC GIẢI")):
        cell = signatures.cell(0, idx)
        cell.width = Cm(8.35)
        p = cell.paragraphs[0]
        style_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER, after=0)
        r1 = p.add_run(title + "\n")
        set_run_font(r1, size=11, bold=True, color=NAVY)
        r2 = p.add_run("(Ký và ghi rõ họ tên)\n\n\n\n")
        set_run_font(r2, size=10, italic=True, color=MUTED)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    line = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:top")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), MID_GRAY)
    line.append(bottom)
    footer_p._p.get_or_add_pPr().append(line)
    run = footer_p.add_run("K-SPORT • Hướng dẫn Check-in VĐV • ")
    set_run_font(run, size=9, color=MUTED)
    add_page_number(footer_p)

    core = doc.core_properties
    core.title = "Hướng dẫn quy trình kiểm tra thẻ và cân VĐV"
    core.subject = "Tài liệu hướng dẫn dành cho HLV/Trưởng đoàn"
    core.author = "Ban Tổ chức Giải"
    core.keywords = "K-SPORT, Karate, Check-in, cân VĐV, kiểm tra thẻ"

    doc.save(OUT_FILE)
    print(OUT_FILE)


if __name__ == "__main__":
    build_document()
